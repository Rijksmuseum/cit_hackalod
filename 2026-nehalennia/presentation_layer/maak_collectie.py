"""Bouwt data/collectie.json: RMO-objecten + 'altaardossiers' met gekoppelde RMA-prenten.

    python3 -I maak_collectie.py pad/naar/edcs_nehalennia_analyse_6.xlsx
    python3 -I maak_collectie.py pad/naar/xlsx --haal-beelden   # ontbrekende iiif-id's ophalen (1 per 1,5 s)

Bronnen
  - data/timeline.json      : RMO-objecten (uit maak_demo_data.py)
  - de EDCS-xlsx            : tabbladen 'Inscripties' en 'Prenten overzicht' (koppeling via EDCS-ID)
  - data/rma_beelden.json   : PID -> iiif.micr.io-id van de RMA-prent

CONCEPT: de koppeling altaar <-> prent komt uit de xlsx (zekerheid/status per prent blijven bewaard).
De koppeling RMO-object <-> EDCS-ID bestaat nog niet; RMO-objecten hebben dus (nog) geen opschrift of prenten.
"""

import json
import re
import sys
import time
import urllib.request
from datetime import date
from pathlib import Path

import openpyxl

HIER = Path(__file__).parent
TIMELINE = HIER / "data" / "timeline.json"
BEELDEN = HIER / "data" / "rma_beelden.json"
UIT = HIER / "data" / "collectie.json"

EDCS_URL = "https://db.edcs.eu/epigr/epi_einzel.php?s_language=en&p_edcs_id={}"  # nog te controleren
RMA_URL = "https://id.rijksmuseum.nl/{}"  # nog te controleren: moet doorverwijzen naar de objectpagina


def norm(edcs):
    return "EDCS-" + re.sub(r"\D", "", edcs).zfill(8)


def leeg(w):
    return None if w is None or str(w).strip() in ("", "-") else str(w).strip()


def lees_xlsx(pad):
    wb = openpyxl.load_workbook(pad, read_only=True)
    ins = {}
    for r in list(wb["Inscripties"].iter_rows(values_only=True))[1:]:
        if r[0]:
            ins[norm(r[0])] = r
    prenten = []
    for r in list(wb["Prenten overzicht"].iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        ids = []
        for deel in str(r[3] or "").split(";"):
            m = re.search(r"EDCS-\d+", deel)
            if m:
                nr = re.search(r"nr\. (\d+)", deel)
                ids.append((norm(m.group()), nr.group(1) if nr else None))
        zek = [z.strip() for z in str(r[4] or "").split(";")]
        prenten.append(dict(obj=r[0], pid=str(r[1]).rstrip("/").split("/")[-1], toont=r[2], ids=ids,
                            zek=zek, status=r[5], opm=leeg(r[6])))
    return ins, prenten


def haal_beelden(prenten, beelden):
    for p in prenten:
        if p["ids"] and p["pid"] not in beelden:
            req = urllib.request.Request(f"https://data.rijksmuseum.nl/{p['pid']}",
                                         headers={"Accept": "application/ld+json"})
            try:
                data = json.load(urllib.request.urlopen(req, timeout=30))
            except Exception as fout:  # noqa: BLE001
                print("mislukt", p["pid"], fout)
                continue
            m = re.search(r"https://iiif\.micr\.io/([^/]+)/", json.dumps(data))
            if m:
                beelden[p["pid"]] = m.group(1)
                print("ok", p["obj"], m.group(1))
            time.sleep(1.5)
    BEELDEN.write_text(json.dumps(beelden, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def main():
    pad = next(a for a in sys.argv[1:] if not a.startswith("--"))
    ins, prenten = lees_xlsx(pad)
    beelden = json.loads(BEELDEN.read_text(encoding="utf-8"))
    if "--haal-beelden" in sys.argv:
        haal_beelden(prenten, beelden)

    per_altaar = {}
    for p in prenten:
        kopie = re.search(r"kopie van ([\w.\-]+)", p["toont"] or "")
        verzamel = len(p["ids"]) > 1 or p["status"] == "deels gekoppeld"
        for i, (edcs, nr) in enumerate(p["ids"]):
            zek = p["zek"][i] if i < len(p["zek"]) else p["zek"][0]
            per_altaar.setdefault(edcs, []).append({
                "objectnummer": p["obj"],
                "beeld": beelden.get(p["pid"]),
                "link": RMA_URL.format(p["pid"]),
                "toont": p["toont"],
                "zekerheid": zek,
                "nr": nr,
                "verzamelprent": verzamel,
                "kopieVan": kopie.group(1) if kopie else None,
                "opmerking": p["opm"],
            })

    def volgorde(pr):  # originelen en zekere koppelingen eerst
        return (pr["kopieVan"] is not None, pr["verzamelprent"], pr["zekerheid"] != "zeker", pr["objectnummer"])

    dossiers = []
    for edcs, lijst in per_altaar.items():
        r = ins[edcs]
        personen = leeg(r[9])
        if not personen:
            titel = "Altaar, naam niet leesbaar"
        elif "[" in personen:
            titel = "Altaar, naam deels leesbaar"
        else:
            titel = "Altaar van " + personen.replace(";", " en")
        dossiers.append({
            "id": edcs,
            "soort": "dossier",
            "objectnaam": "altaar",
            "titel": titel,
            "instelling": "RMA",   # van de gekoppelde prenten; de instelling van het altaar zelf is niet bekend
            "vindplaats": leeg(r[3]) and r[3].split("/")[0].strip(),
            "datering": leeg(r[6]),
            "opschrift": [regel.strip() for regel in str(r[5]).split(" / ")],
            "vertaling": leeg(r[8]),
            "opschriftLink": EDCS_URL.format(edcs),
            "edcsId": edcs,
            "prenten": sorted(lijst, key=volgorde),
        })
    dossiers.sort(key=lambda d: (-len(d["prenten"]), d["id"]))

    rmo = []
    for jaar in json.loads(TIMELINE.read_text(encoding="utf-8"))["jaren"]:
        for o in jaar["objecten"]:
            rmo.append({
                "id": o["id"], "soort": "rmo", "objectnaam": o["objectnaam"], "titel": o["titel"],
                "instelling": "RMO", "materiaal": o["materiaal"], "vindplaats": o["vindplaats"],
                "afbeelding": o["afbeelding"], "bronLink": o["bronLink"],
                "objectnummer": None,   # inventarisnummer zit niet in de huidige oogst
                "recordnummer": o["id"],
            })

    UIT.write_text(json.dumps({
        "gegenereerd": date.today().isoformat(),
        "concept": True,
        "stukken": dossiers + rmo,
    }, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"{len(dossiers)} altaardossiers ({sum(len(d['prenten']) for d in dossiers)} prentkoppelingen) + {len(rmo)} RMO-objecten -> {UIT}")


if __name__ == "__main__":
    main()
